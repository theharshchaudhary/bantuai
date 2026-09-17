"""R2 tests: the knowledge folder - indexing, syncing, searching, tools, settings.

No API key, no network. Everything happens in temp folders, never the user's
real Documents\\Bantu Knowledge.

    .venv/Scripts/python.exe tests/test_r2_knowledge.py
"""

from __future__ import annotations

import os
import sys
import tempfile
import time
from pathlib import Path

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
os.environ["APPDATA"] = tempfile.mkdtemp(prefix="bantu_knowledge_appdata_")
ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

for _s in (sys.stdout, sys.stderr):
    try:
        _s.reconfigure(encoding="utf-8", errors="replace")
    except (AttributeError, ValueError):
        pass

import core.knowledge as knowledge_module
from core.knowledge import CHUNK_CHARS, Knowledge, chunk_text, read_text_file
from core.knowledge import register as register_knowledge
from core.memory import Memory
from core.providers.base import ToolCall
from core.tools.registry import Tier, ToolRegistry

PASS, FAIL = [], []


def check(name, cond, detail=""):
    (PASS if cond else FAIL).append(name)
    print(("  ok   " if cond else "  FAIL ") + name + (f"  -> {detail}" if detail and not cond else ""))


def fresh(readers=None):
    base = Path(tempfile.mkdtemp())
    memory = Memory(base / "k.db")
    return memory, Knowledge(memory.db, base / "Bantu Knowledge", readers), base / "Bantu Knowledge"


def touch_later(path: Path) -> None:
    """Make a rewrite visibly newer, whatever the filesystem's timestamp resolution."""
    st = path.stat()
    os.utime(path, (st.st_atime + 5, st.st_mtime + 5))


# --- chunking and reading ---------------------------------------------------------

print("\n[chunking and reading]")
check("nothing to chunk gives no chunks", chunk_text("") == [] and chunk_text("\n\n  \n") == [])
paras = "\n\n".join(f"Paragraph {i}. " + "Some policy detail here. " * 8 for i in range(12))
chunks = chunk_text(paras)
check("chunks stay within the size", all(len(c) <= CHUNK_CHARS for c in chunks), str([len(c) for c in chunks]))
check("small paragraphs are packed together, not one chunk each", 1 < len(chunks) < 12, str(len(chunks)))
check("no text is lost", " ".join(" ".join(chunks).split()) == " ".join(paras.split()))
long_para = " ".join(f"Sentence number {i} explains one more rule." for i in range(80))
check("a long paragraph splits at sentences", all(len(c) <= CHUNK_CHARS and c.rstrip().endswith(".")
                                                   for c in chunk_text(long_para)))
check("one enormous run-on is cut to size", all(len(c) <= CHUNK_CHARS for c in chunk_text("x" * 3500)))

tmp = Path(tempfile.mkdtemp())
(tmp / "utf8.txt").write_text("नमस्ते café", encoding="utf-8")
(tmp / "bom.txt").write_text("with a BOM", encoding="utf-8-sig")
(tmp / "wide.txt").write_text("saved as UTF-16 by Notepad", encoding="utf-16")
(tmp / "old.txt").write_bytes("caf\xe9 in Windows-1252".encode("cp1252"))
check("UTF-8 text reads, Devanagari included", read_text_file(tmp / "utf8.txt") == "नमस्ते café")
check("a UTF-8 byte-order mark is dropped", read_text_file(tmp / "bom.txt") == "with a BOM")
check("UTF-16 (Notepad's 'Unicode') reads", read_text_file(tmp / "wide.txt") == "saved as UTF-16 by Notepad")
check("old Windows-1252 text still reads", read_text_file(tmp / "old.txt") == "café in Windows-1252")


# --- syncing ----------------------------------------------------------------------

print("\n[syncing the folder]")
memory, kb, folder = fresh()
check("the folder is created on first sync", not folder.exists() and kb.sync() is not None and folder.is_dir())
(folder / "leave_policy.md").write_text(
    "# Leave\n\nEmployees receive 18 days of paid leave per year.\n\nUp to 5 unused days carry over.",
    encoding="utf-8")
(folder / "sub").mkdir()
(folder / "sub" / "vendors.txt").write_text("Globex is the preferred vendor for dashboards.", encoding="utf-8")
(folder / "contacts.csv").write_text("name,role\nRam,vendor contact\n", encoding="utf-8")
(folder / "नीति.txt").write_text("कर्मचारियों को हर साल अठारह दिन की सवेतन छुट्टी मिलती है।", encoding="utf-8")
(folder / "setup.exe").write_bytes(b"MZ")
(folder / "photo.png").write_bytes(b"\x89PNG")
(folder / ".hidden.txt").write_text("secret scratch", encoding="utf-8")
(folder / "~$leave_policy.md").write_text("Word lock file", encoding="utf-8")

report = kb.sync()
check("supported files are indexed, subfolders included",
      sorted(report.added) == ["contacts.csv", "leave_policy.md", "sub/vendors.txt", "नीति.txt"], str(report.added))
check("programs, images, hidden files and Office lock files are ignored",
      not any(n in " ".join(report.added) for n in ("setup", "photo", "hidden", "~$")))
check("a second sync with nothing changed does nothing", not kb.sync().changed)

hits = kb.search("paid leave days")
check("search finds the passage and names the file", hits and hits[0]["path"] == "leave_policy.md"
      and "18 days" in hits[0]["text"], str(hits))
check("a file in a subfolder is found by its relative path", kb.search("preferred vendor")[0]["path"] == "sub/vendors.txt")
check("a Hindi document is found by a Hindi word", kb.search("छुट्टी") and kb.search("छुट्टी")[0]["path"] == "नीति.txt")
check("nothing matching returns nothing", kb.search("parking spaces") == [])

(folder / "leave_policy.md").write_text("Employees receive 21 days of paid leave per year.", encoding="utf-8")
touch_later(folder / "leave_policy.md")
report = kb.sync()
check("an edited file is re-read", report.updated == ["leave_policy.md"], str(report))
check("...its new text is found and the old text is gone",
      "21 days" in kb.search("paid leave")[0]["text"] and not any("18 days" in h["text"] for h in kb.search("paid leave")))
check("...and the old 'carry over' passage no longer matches", kb.search("carry over") == [])

(folder / "sub" / "vendors.txt").unlink()
report = kb.sync()
check("a deleted file is dropped from the index", report.removed == ["sub/vendors.txt"] and kb.search("Globex") == [],
      str(report))
check("the summary counts what is indexed", kb.summary() == "3 documents indexed", kb.summary())


def broken_reader(path):
    raise ValueError("encrypted PDF")


memory, kb, folder = fresh({".pdf": broken_reader})
folder.mkdir(parents=True)
(folder / "locked.pdf").write_bytes(b"%PDF-1.7")
(folder / "blank.txt").write_text("   \n\n  ", encoding="utf-8")
(folder / "fine.txt").write_text("The office opens at nine.", encoding="utf-8")
report = kb.sync()
failed = dict(report.failed)
check("a file that cannot be read is reported, and the rest still index",
      "encrypted PDF" in failed.get("locked.pdf", "") and report.added == ["fine.txt"], str(report))
check("a file with no text is reported as having none", "no readable text" in failed.get("blank.txt", ""), str(failed))
check("a failed file is not re-read on every poll", not kb.sync().changed)
check("the summary says some could not be read", kb.summary() == "1 document indexed, 2 could not be read", kb.summary())
check("the file list shows why", any(f["path"] == "locked.pdf" and "encrypted" in (f["error"] or "") for f in kb.files()))

original_max = knowledge_module.MAX_FILE_BYTES
knowledge_module.MAX_FILE_BYTES = 10
memory, kb, folder = fresh()
folder.mkdir(parents=True)
(folder / "huge.txt").write_text("far more than ten bytes of text", encoding="utf-8")
report = kb.sync()
knowledge_module.MAX_FILE_BYTES = original_max
check("an oversized file is skipped with a reason, not read", "larger than" in dict(report.failed).get("huge.txt", ""),
      str(report))


# --- real document readers ---------------------------------------------------------

print("\n[Word and Excel documents]")
from platform_desktop.files import extract_document_text

readers = {".docx": extract_document_text, ".xlsx": extract_document_text}
memory, kb, folder = fresh(readers)
folder.mkdir(parents=True)
import docx
import openpyxl

doc = docx.Document()
doc.add_heading("Travel policy", 1)
doc.add_paragraph("Economy class for flights under six hours.")
doc.save(str(folder / "travel.docx"))
wb = openpyxl.Workbook()
wb.active.append(["Vendor", "Contact"])
wb.active.append(["Globex", "Ram Sharma"])
wb.save(str(folder / "vendors.xlsx"))
report = kb.sync()
check("Word and Excel files index through the same readers as read_document",
      sorted(report.added) == ["travel.docx", "vendors.xlsx"], str(report))
check("a Word document is searchable", kb.search("economy flights")[0]["path"] == "travel.docx")
check("a spreadsheet row is searchable", kb.search("Ram Sharma")[0]["path"] == "vendors.xlsx")


# --- polling ------------------------------------------------------------------------

print("\n[polling]")
memory, kb, folder = fresh()
kb.start_polling(every=0.2)
time.sleep(0.3)
(folder / "late.txt").write_text("Dropped in while Bantu was running.", encoding="utf-8")
deadline = time.time() + 5
while time.time() < deadline and not kb.search("running"):
    time.sleep(0.1)
check("a file dropped in while running is picked up by the poll", bool(kb.search("running")))
kb.stop_polling()
check("polling stops", not (kb._poller and kb._poller.is_alive()))


# --- tools ----------------------------------------------------------------------------

print("\n[knowledge tools]")
memory, kb, folder = fresh()
reg = ToolRegistry()
register_knowledge(reg, kb)
check("searching and listing run without asking",
      reg.tools["search_knowledge"].tier is Tier.AUTO and reg.tools["list_knowledge"].tier is Tier.AUTO)
empty = reg.execute(ToolCall("l", "list_knowledge", {}))
check("an empty folder says where it is and what it accepts", "empty" in empty and str(folder) in empty
      and ".md" in empty, empty)
(folder / "wifi.txt").write_text("The guest wifi password is written on the fridge.", encoding="utf-8")
found = reg.execute(ToolCall("s", "search_knowledge", {"query": "guest wifi"}))
check("search syncs first, so a file dropped a moment ago is found", "From wifi.txt (part 1)" in found, found)
partial = reg.execute(ToolCall("s", "search_knowledge", {"query": "guest parking"}))
check("a passage sharing only some words is labelled a near miss, not an answer",
      partial.startswith("No passage mentions all of") and "wifi.txt" in partial, partial)
check("a passage with every word is returned plainly",
      not reg.execute(ToolCall("s", "search_knowledge", {"query": "wifi password"})).startswith("No passage"))
none = reg.execute(ToolCall("s", "search_knowledge", {"query": "parking"}))
check("finding nothing says the documents do not cover it", "do not cover it" in none and "1 document indexed" in none,
      none)
listed = reg.execute(ToolCall("l", "list_knowledge", {}))
check("list_knowledge lists files", "wifi.txt: 1 part(s)" in listed, listed)

lazy = ToolRegistry()
register_knowledge(lazy, kb)
lazy.enable_lazy_loading(base={"core"})
catalog = next(s for s in lazy.specs() if s.name == "load_tools").description
check("knowledge tools load on demand, with a catalog line saying what they are for",
      "search_knowledge" not in [s.name for s in lazy.specs()] and "knowledge:" in catalog
      and "what their documents or notes say" in catalog, catalog)


# --- where the folder is, and Settings ------------------------------------------------

print("\n[folder location and Settings]")
import main
from core import config as cfg
from platform_desktop.paths import documents_dir

default = main.knowledge_folder(cfg.Settings())
check("by default the folder is Documents\\Bantu Knowledge, where people can find it",
      default == documents_dir() / "Bantu Knowledge", str(default))
check("...not inside Bantu's own data folder, which tools may not touch",
      cfg.user_data_dir() not in default.parents)
check("a configured folder is used instead", main.knowledge_folder(cfg.Settings(knowledge_dir=str(folder))) == folder)
check("the settings default leaves it unset", cfg.Settings().knowledge_dir == "")

from PyQt5.QtWidgets import QApplication

app = QApplication.instance() or QApplication([])
from ui.settings_dialog import SettingsDialog
from ui.setup_parts import Services

services = Services(check_key=lambda p, k: None, store_key=lambda p, k: None,
                    existing_key=lambda p: (None, None), list_devices=lambda: [],
                    preview_voice=lambda s, g, t: None, open_url=lambda u: None)
dlg = SettingsDialog(cfg.Settings(), services, data_dir=str(folder.parent), knowledge=kb)
check("Settings shows the knowledge folder and what is indexed",
      dlg.knowledge_summary.text() == "1 document indexed", dlg.knowledge_summary.text())
plain = SettingsDialog(cfg.Settings(), services)
check("Settings still opens without a knowledge folder", not hasattr(plain, "knowledge_summary"))

print(f"\n{len(PASS)} passed, {len(FAIL)} failed")
for f in FAIL:
    print("  FAILED:", f)
sys.exit(1 if FAIL else 0)
